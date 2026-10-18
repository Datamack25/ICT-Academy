"""Génère le document Word (.docx) des scripts TradingView avec python-docx."""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .pine import SCRIPTS

BLUE = RGBColor(0x1F, 0x4E, 0x9A)
BLACK = RGBColor(0, 0, 0)

READ_NUMBERS = [
    ("avg RR (gain moyen en R)", "Résultat net en R divisé par le nombre de trades. Exemple : 10 trades à 1:2, 4 gagnants et 6 perdants "
     "→ 4 × 2 − 6 = +2 R → avg RR = 2 / 10 = +0,20. Sous 0 : la stratégie perd de l'argent, même avec 99 % de réussite. "
     "+0,05 à +0,10 : petit avantage ; +0,20 : bon avantage ; +0,40 ou plus : exceptionnel et rare, vérifier qu'il n'y a pas d'erreur de backtest."),
    ("PF (profit factor)", "Total gagné ÷ total perdu. 8 $ / 6 $ = 1,33 : pour 1 $ perdu, 1,33 $ gagné. Sous 1 : perte ; 1,0–1,2 : minimum, "
     "quelques frais suffisent à basculer en perte ; 1,3–2,0 : bon ; au-delà de 3 : trop beau, vérifier le backtest."),
    ("t-score", "Le résultat est-il de la chance ? t = avg RR ÷ (dispersion des résultats ÷ √trades). Sous 2 : stratégie non prouvée ; "
     "2 ou plus : avantage probable mais à re-tester ; 3 ou plus : avantage solide (seuil exigé pour un combo choisi dans une grille)."),
    ("Max DD (drawdown maximal)", "La plus grosse baisse depuis un sommet. Exemple : +10 R puis redescente à +4 R = 6 trades perdus → Max DD = 6 R. "
     "C'est lui qui décide si un compte prop survit."),
    ("Pertes consécutives", "Sur 100 trades à 40 % de réussite, 9 pertes d'affilée sont possibles : la taille de position doit permettre d'y survivre."),
    ("Win rate seul", "Ne veut rien dire sans le ratio gain/perte : 72 % de réussite avec une cible minuscule et un stop large peut perdre "
     "de l'argent après coûts. Toujours lire avg RR et PF APRÈS coûts (1 point par trade)."),
]


def _shade(paragraph, hex_fill="F2F4F8"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def _cell_shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, h in enumerate(header):
        c = t.rows[0].cells[k]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0xFF, 0xFF, 0xFF)
        _cell_shade(c, "1F4E9A")
    for row in rows:
        cells = t.add_row().cells
        for k, v in enumerate(row):
            cells[k].text = ""
            r = cells[k].paragraphs[0].add_run(str(v))
            r.font.size, r.font.color.rgb = Pt(9), BLACK
    if widths:
        for row in t.rows:
            for k, w in enumerate(widths):
                row.cells[k].width = Cm(w)
    return t


def build_docx(verdicts: list | None = None, data_label: str = "") -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.font.color.rgb = BLACK
    for lvl in (1, 2, 3):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.color.rgb = BLUE
        hs.font.name = "Calibri"

    title = doc.add_heading("Scripts TradingView — ICT × MMM pour le MNQ", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.add_run(f"ICT MNQ Academy · généré le {datetime.now():%d/%m/%Y %H:%M}").italic = True
    doc.add_paragraph(
        "Ce document regroupe six scripts Pine Script v6 qui reproduisent dans TradingView les dessins et les règles de la plateforme : "
        "la méthodologie ICT du cours Warren (structure, orderblocks, breakers, FVG, liquidité, OTE, kill zones, Monday range) et la "
        "lecture « order flow » inspirée du Max Maserati Model (boussole des bougies, GMC, escalier T#, boîtes de session, MMPD, "
        "alignement multi-unités de temps). « Edge 5 Desk » est le poste de pilotage : il simule les 5 setups à la clôture, trace les "
        "trades M#T# avec leur résultat et affiche les tableaux EDGE DESK et PROP HELPER, avec exactement les mêmes règles que la page "
        "Edge Desk de la plateforme. La stratégie « Edge Lab » permet de backtester ces setups dans le Strategy Tester avec "
        "les mêmes statistiques que la plateforme (avg RR après coûts, PF, t-score, drawdown en R, verdict).")

    doc.add_heading("Installer un script", 1)
    for k, step in enumerate([
        "Ouvre TradingView sur MNQ1! (CME) en 15 minutes — l'unité de décision recommandée.",
        "Clique sur « Pine Editor » en bas de l'écran, puis « Ouvrir » → « Nouvel indicateur » (ou « Nouvelle stratégie » pour Edge Lab).",
        "Efface tout le contenu, colle le script complet (de //@version=6 jusqu'à la dernière ligne).",
        "Clique sur « Enregistrer », donne un nom, puis « Ajouter au graphique ».",
        "Ouvre les paramètres (roue dentée) pour régler la sensibilité des swings, les sessions et les filtres.",
        "Pour Edge Lab : onglet « Strategy Tester » → onglet « Performance » ; le tableau en bas à droite du graphique donne avg RR, t-score et verdict.",
        "Pour Edge 5 Desk : clic droit sur le graphique → « Ajouter une alerte » → condition « Edge 5 Desk » → « Any alert() function call » "
        "→ fréquence « Once per bar close » : chaque trade t'envoie entrée, SL, ½ TP et TP.",
        "Charge le plus d'historique possible (fais défiler le graphique vers la gauche) : les verdicts et le prop helper sont calculés sur "
        "les bougies chargées ; en dessous de 30 trades par setup, un verdict n'est pas fiable.",
    ], 1):
        doc.add_paragraph(f"{k}. {step}")
    doc.add_paragraph("Si TradingView affiche une erreur de compilation, copie le message et la ligne indiquée : les scripts ont été écrits pour "
                      "Pine v6 et relus ligne à ligne, mais ils n'ont pas pu être compilés dans TradingView au moment de la génération.")

    doc.add_heading("Avant de trader en réel", 1)
    for line in [
        "N'exécute en réel que les setups au statut TRADE sur ton propre historique (6 à 12 mois en 15 min), après au moins 30 trades "
        "en paper / forward-test cohérents avec le backtest.",
        "Risque 1 % max par trade (0,5 % les jours d'annonces), stop posé avant l'entrée, 2 positions max, arrêt de la journée à −3 %.",
        "Évite les jours CPI, NFP et FOMC (le site Max Maserati mesure −0,04 R en moyenne ces jours-là).",
        "Un backtest ne garantit pas l'avenir : ces scripts sont des outils d'analyse et d'entraînement, pas un conseil en investissement.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Vue d'ensemble", 1)
    _table(doc, ["Script", "Type", "Rôle"], [(n, k, d) for n, f, k, c, d in SCRIPTS], widths=[4.2, 3.2, 10])

    if verdicts:
        doc.add_heading("Verdicts du dernier backtest de la plateforme", 1)
        doc.add_paragraph(f"Données : {data_label}. Coût : 1 point par trade. Période ancienne = 70 % des dates, récente = 30 %.")
        _table(doc, ["Statut", "Setup", "Trades", "WR", "avg RR ancien → récent", "PF", "t-score"],
               [(v["verdict"], v["setup"], v["n"], v["wr"], v["old_new"], v["pf"], v["t"]) for v in verdicts],
               widths=[2.6, 5.6, 1.5, 1.4, 3.2, 1.3, 1.6])
        doc.add_paragraph("Lecture des statuts : TRADE = positif dans les deux périodes avec un t-score suffisant ; USE AS FILTER = n'est pas "
                          "un trade seul mais améliore les setups auxquels on l'ajoute ; WATCH = positif dans les deux périodes mais trop peu de "
                          "trades ou t-score < 2, à tester en avant ; DON'T TRADE = perd, s'est effacé, ou ne marche que sur une période.")

    doc.add_heading("Comment lire les chiffres", 1)
    for k, v in READ_NUMBERS:
        para = doc.add_paragraph()
        para.add_run(k + " — ").bold = True
        para.add_run(v)

    doc.add_heading("Règles d'exécution retenues", 1)
    for rule in [
        "Entrer à la clôture de la bougie signal : les limites profondes sont les pires entrées (un remplissage profond signale souvent une cassure ratée).",
        "Stops basés sur la volatilité (ATR) plutôt que des stops de structure serrés, qui se font toucher par le bruit.",
        "Prendre la moitié à +0,5 R puis remonter le stop à l'entrée : c'est ce qui fait passer le win rate au-dessus de 60 %.",
        "Trader l'unité 15 min pendant l'Asie (sauf 00h–02h) et New York matin (09h30–13h30, heure NY). Éviter 1–5 min : les coûts mangent l'avantage.",
        "Utiliser le MMPD comme gate de momentum (long en XP, short en XD) — ne jamais jouer le retournement des extrêmes.",
        "Éviter les journées CPI, NFP et FOMC en entier. Maximum 2 positions ouvertes.",
        "ICT : ne prendre un breaker ou un OB qu'après un MSS validé à la clôture, idéalement précédé d'une prise de liquidité (méthode C).",
    ]:
        doc.add_paragraph(rule, style="List Bullet")

    for n, fname, kind, code, desc in SCRIPTS:
        doc.add_page_break()
        doc.add_heading(n, 1)
        meta = doc.add_paragraph()
        meta.add_run("Fichier : ").bold = True
        meta.add_run(f"{fname} · {kind}")
        doc.add_paragraph(desc)
        doc.add_heading("Code Pine Script v6", 2)
        for line in code.strip("\n").split("\n"):
            cp = doc.add_paragraph()
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.line_spacing = 1.0
            run = cp.add_run(line if line else " ")
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(7.5)
            run.font.color.rgb = BLACK
            _shade(cp)

    doc.add_page_break()
    doc.add_heading("Avertissement", 1)
    doc.add_paragraph("Outil pédagogique d'entraînement. Les résultats de backtest ne préjugent pas des résultats futurs ; aucune "
                      "information de ce document n'est un conseil en investissement. Teste toujours un script sur compte de démonstration.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
